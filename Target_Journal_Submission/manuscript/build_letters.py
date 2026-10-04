"""Cover letter and highlights (DOCX). Numbers are read from the result files."""
import json, os
import pandas as pd
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

HERE = os.path.dirname(os.path.abspath(__file__))
R = os.path.join(os.path.dirname(HERE), "code", "results")
s = pd.read_csv(os.path.join(R, "summary_results.csv")); s = s[s.alpha == 0.1]
g = lambda t, m, c="coverage": float(s[(s.task == t) & (s.method == m)][c].iloc[0])
TITLE = ("Group-conformal calibration of near-infrared prediction intervals for new seasons, instruments and regions "
         "using PLS diagnostics")


def doc():
    d = Document()
    st = d.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(12)
    for sec in d.sections:
        sec.left_margin = sec.right_margin = Cm(2.5); sec.top_margin = sec.bottom_margin = Cm(2.5)
    return d


def cover_letter(path):
    d = doc()
    for line in ["[Corresponding author name]", "[Department, Institution]", "[Address]", "[E-mail]", "", "[Date]", "",
                 "The Editors", "Talanta", "Elsevier", ""]:
        d.add_paragraph(line)
    d.add_paragraph("Dear Editors,")
    paras = [
        f"We submit the manuscript “{TITLE}” for consideration as a Research Paper in Talanta.",
        "Near-infrared calibrations are routinely deployed on harvest seasons, instruments and regions that were not part "
        "of the calibration set, yet their prediction intervals are calibrated as if future samples resembled the "
        "calibration samples. Using four public data sets (mango dry matter, soil organic carbon, corn and pharmaceutical "
        "tablets; 146,873 spectra) and twelve deployment-shift scenarios with sample-disjoint group splits, we show that "
        "classical PLS intervals and randomly calibrated conformal intervals under-cover such new groups — for example, "
        f"random-split conformal intervals covered {g('mango_season_forward','SCP'):.2f} of next-season mango spectra and "
        f"{g('mango_instrument','SCP'):.2f} of new-instrument spectra at a nominal 0.90.",
        "We propose group-conformal calibration (GC-D): conformity scores are computed out-of-group with group-balanced "
        "weights, re-using the group-wise cross-validation chemometricians already perform, and are scaled by the "
        "Hotelling T² and Q diagnostics so that intervals widen for atypical spectra. GC-D covered "
        f"{g('mango_season_forward','GC-D'):.2f}–{g('mango_season_loso','GC-D'):.2f} for new seasons, "
        f"{g('mango_instrument','GC-D'):.2f} for new instruments and {g('ossl_lucas_block','GC-D'):.2f} for new soil regions, "
        "with interval scores close to an oracle calibrated on the new group, kept its coverage under simulated wavelength "
        "drift, and reduced false acceptances in specification-limit decisions. We also report its limits candidly: formal "
        "group-level guarantees need about ten or more calibration groups, a classical interval based on group-wise RMSECV "
        "captures much of the benefit, and no interval method can compensate for an uncorrected instrument transfer.",
        "We believe the work fits the scope of Talanta because it addresses a core analytical question — how to state the "
        "uncertainty of a multivariate calibration result reliably when the method is applied in practice — and evaluates "
        "it on real analytical data with the consequences for analytical decisions. The experimental protocol and "
        "hypotheses were fixed before the held-out evaluation (additional and post-hoc analyses are labelled as such), all data "
        "are public, and the complete code regenerates every number, table and figure.",
        "This manuscript is original, has not been published previously, and is not under consideration for publication "
        "elsewhere. All authors have approved the submission and declare no competing interests. No human or animal "
        "subjects were involved. The use of AI-assisted tools is declared in the manuscript.",
        "[Optional: suggested reviewers with expertise in chemometrics and conformal prediction — to be completed by the authors.]",
        "Thank you for considering our manuscript.",
    ]
    for p in paras:
        para = d.add_paragraph(p); para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    d.add_paragraph(""); d.add_paragraph("Yours sincerely,"); d.add_paragraph(""); d.add_paragraph("[Corresponding author], on behalf of all authors")
    d.save(path)


def highlights(path):
    H = json.load(open(os.path.join(HERE, "highlights.json")))
    assert all(len(h) <= 85 for h in H) and 3 <= len(H) <= 5
    d = doc(); d.add_heading("Highlights", level=1)
    for h in H:
        d.add_paragraph(h, style="List Bullet")
    d.save(path)


if __name__ == "__main__":
    cover_letter(os.path.join(HERE, "cover_letter.docx"))
    highlights(os.path.join(HERE, "highlights.docx"))
    print("written cover_letter.docx, highlights.docx")
